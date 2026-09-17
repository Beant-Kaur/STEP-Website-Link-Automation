import os
import openpyxl
import pytest

from database.models import DatabaseRepository, LinkRecord
from pipeline.orchestrator import PipelineOrchestrator
from reference import (
    ExcelReferenceDatabase,
    RegulatoryDatabaseError,
    RegulatoryResource,
    RegulatoryAuthority,
    ResourceVersion,
)


class TestExcelReferenceDatabaseLoading:
    def test_load_real_workbook(self):
        db = ExcelReferenceDatabase()
        db.load()
        assert db.loaded is True
        assert len(db.resources) >= 3
        assert len(db.authorities) >= 4

        # Check India pilot records
        ind_001 = db.get_resource("IND-001")
        assert ind_001 is not None
        assert ind_001.country == "India"
        assert "SEBI" in ind_001.name or "BRSR" in ind_001.name
        assert ind_001.official_domain == "sebi.gov.in"
        assert ind_001.issuing_authority == "Securities and Exchange Board of India (SEBI)"
        assert len(ind_001.versions) >= 2
        assert "Do not mark as replaced solely from age" in ind_001.validator_rule

        ind_002 = db.get_resource("IND-002")
        assert ind_002 is not None
        assert ind_002.country == "India"
        assert "Companies Act" in ind_002.name
        assert "Verify current consolidated text" in ind_002.validator_rule

        ind_003 = db.get_resource("IND-003")
        assert ind_003 is not None
        assert ind_003.country == "India"
        assert "RBI" in ind_003.name or "Reserve Bank" in ind_003.name
        assert "Human review" in ind_003.validator_rule

    def test_authorities_parsed_correctly(self):
        db = ExcelReferenceDatabase()
        db.load()
        assert "AUTH-IN-001" in db.authorities
        auth1 = db.authorities["AUTH-IN-001"]
        assert auth1.authority_name == "Securities and Exchange Board of India (SEBI)"
        assert auth1.official_domain == "sebi.gov.in"
        assert auth1.trust_level == "Official"

    def test_to_dict_conversion(self):
        db = ExcelReferenceDatabase()
        db.load()
        ind_001 = db.get_resource("IND-001")
        data = ind_001.to_dict()
        assert isinstance(data, dict)
        assert data["resource_id"] == "IND-001"
        assert data["country"] == "India"
        assert isinstance(data["versions"], list)
        assert len(data["versions"]) > 0
        assert data["authority_details"]["official_domain"] == "sebi.gov.in"


class TestExcelErrorHandling:
    def test_missing_file_raises_clear_error(self, tmp_path):
        non_existent = str(tmp_path / "does_not_exist.xlsx")
        db = ExcelReferenceDatabase(file_path=non_existent)
        with pytest.raises(RegulatoryDatabaseError) as exc_info:
            db.load()
        assert "not found" in str(exc_info.value).lower()

    def test_missing_sheet_raises_clear_error(self, tmp_path):
        invalid_wb_path = str(tmp_path / "invalid_sheets.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Resources"
        wb.save(invalid_wb_path)
        wb.close()

        db = ExcelReferenceDatabase(file_path=invalid_wb_path)
        with pytest.raises(RegulatoryDatabaseError) as exc_info:
            db.load()
        assert "missing from Excel file" in str(exc_info.value)
        assert "Authorities" in str(exc_info.value)

    def test_missing_column_raises_clear_error(self, tmp_path):
        invalid_wb_path = str(tmp_path / "invalid_columns.xlsx")
        wb = openpyxl.Workbook()
        ws_res = wb.active
        ws_res.title = "Resources"
        ws_res.append(["Resource ID", "Country"])  # Missing required columns
        ws_auth = wb.create_sheet("Authorities")
        ws_auth.append(["Authority ID", "Country", "Authority", "Official Domain"])
        ws_ver = wb.create_sheet("Resource_Versions")
        ws_ver.append(["Resource ID", "Resource / Update", "Relationship"])
        wb.save(invalid_wb_path)
        wb.close()

        db = ExcelReferenceDatabase(file_path=invalid_wb_path)
        with pytest.raises(RegulatoryDatabaseError) as exc_info:
            db.load()
        assert "Required column" in str(exc_info.value)
        assert "missing in sheet 'Resources'" in str(exc_info.value)


class TestIndiaResourceMatching:
    def setup_method(self):
        self.db = ExcelReferenceDatabase()
        self.db.load()

    def test_match_sebi_brsr(self):
        url = "https://www.sebi.gov.in/legal/circulars/may-2021/business-responsibility-and-sustainability-reporting-by-listed-entities_50096.html"
        text = "Business Responsibility and Sustainability Reporting (BRSR)"
        matched = self.db.match_resource(jurisdiction="India", url=url, text=text)
        assert matched is not None
        assert matched.resource_id == "IND-001"
        assert matched.official_domain == "sebi.gov.in"

    def test_match_mca_companies_act(self):
        url = "https://www.mca.gov.in/Ministry/pdf/CompaniesAct2013.pdf"
        text = "Companies Act 2013 (see Section 135 & CSR Rules)"
        matched = self.db.match_resource(jurisdiction="India", url=url, text=text)
        assert matched is not None
        assert matched.resource_id == "IND-002"
        assert "Companies Act" in matched.name

    def test_match_rbi_guidelines(self):
        url = "https://www.rbi.org.in/Scripts/NotificationUser.aspx?Id=12213&Mode=0"
        text = "Reserve Bank of India Guidelines on Environmental and Social Risk Management"
        matched = self.db.match_resource(jurisdiction="India", url=url, text=text)
        assert matched is not None
        assert matched.resource_id == "IND-003"
        assert "RBI" in matched.name or "Reserve Bank" in matched.name


class TestOrchestratorReferenceEnrichment:
    def test_orchestrator_enriches_matched_india_record(self, tmp_path):
        db_path = str(tmp_path / "test_pipeline.db")
        repo = DatabaseRepository(db_path)
        orchestrator = PipelineOrchestrator(repository=repo)

        # Mock HTML containing an India STEP link
        sample_html = """
        <div class="accordion-item">
            <div class="accordion-title">India</div>
            <div class="accordion-collapse">
                <p><strong>Business Responsibility and Sustainability Reporting:</strong>
                <a href="https://www.sebi.gov.in/legal/circulars/may-2021/business-responsibility-and-sustainability-reporting-by-listed-entities_50096.html">
                    SEBI BRSR Circular
                </a></p>
            </div>
        </div>
        """

        # Run with limit=1 and offline mock check to verify enrichment without hitting network
        orchestrator.extractor.extract_from_html = lambda html, base_url, section="": [{
            "url": "https://www.sebi.gov.in/legal/circulars/may-2021/business-responsibility-and-sustainability-reporting-by-listed-entities_50096.html",
            "jurisdiction": "India",
            "step_description": "Business Responsibility and Sustainability Reporting",
            "anchor_text": "SEBI BRSR Circular",
            "section_heading": "India",
        }]

        # Mock URL checker to return clean 200 without network requests
        orchestrator.url_checker.check = lambda url: {
            "http_status": 200,
            "technical_status": "ACCESSIBLE",
            "access_status": "ACCESSIBLE",
            "final_url": url,
            "content_type": "text/html",
            "page_title": "SEBI BRSR Circular 2021",
        }
        orchestrator._fetch_metadata = lambda ct, url, check=None: {
            "title": "SEBI BRSR Circular 2021",
            "access_status": "ACCESSIBLE",
            "technical_status": "ACCESSIBLE",
            "document_type": "Circular",
            "text_sample": "Business Responsibility and Sustainability Reporting by listed entities",
        }

        run_id = orchestrator.run(sample_html, limit=1)
        records = repo.get_all_links()
        assert len(records) >= 1

        rec = records[-1]
        assert rec.jurisdiction == "India"
        assert rec.source_organisation == "Securities and Exchange Board of India (SEBI)"
        assert rec.authority_status == "OFFICIAL_REGULATORY_SOURCE"
        assert "IND-001" in rec.evidence
        assert "sebi.gov.in" in rec.evidence
