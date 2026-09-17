import glob
import logging
import os
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

import openpyxl

from .models import RegulatoryAuthority, RegulatoryResource, ResourceVersion

logger = logging.getLogger(__name__)


class RegulatoryDatabaseError(Exception):
    """Raised when the regulatory reference database Excel file is missing, corrupt, or invalid."""
    pass


class ExcelReferenceDatabase:
    """
    Excel Reference Database Importer & Query Engine.
    Acts as the source-of-truth reference layer for regulatory validations.
    """

    REQUIRED_SHEETS = ["Resources", "Authorities", "Resource_Versions"]

    REQUIRED_COLUMNS = {
        "Resources": [
            "Resource ID",
            "Country",
            "Resource Name / Title",
            "Issuing Authority",
            "Official Domain",
            "Official URL",
        ],
        "Authorities": [
            "Authority ID",
            "Country",
            "Authority",
            "Official Domain",
        ],
        "Resource_Versions": [
            "Resource ID",
            "Resource / Update",
            "Relationship",
        ],
    }

    def __init__(self, file_path: Optional[str] = None):
        self.file_path: Optional[str] = file_path
        self.authorities: Dict[str, RegulatoryAuthority] = {}      # authority_id -> RegulatoryAuthority
        self.resources: Dict[str, RegulatoryResource] = {}          # resource_id -> RegulatoryResource
        self.resources_by_country: Dict[str, List[RegulatoryResource]] = {}
        self.loaded: bool = False

    def find_default_file(self) -> Optional[str]:
        """Locates the Excel database in workspace root or project folders."""
        candidates = [
            self.file_path,
            "India_Regulatory_Reference_Database_Refined (1).xlsx",
            os.path.join(os.path.dirname(__file__), "..", "India_Regulatory_Reference_Database_Refined (1).xlsx"),
            os.path.join(os.path.dirname(__file__), "..", "..", "India_Regulatory_Reference_Database_Refined (1).xlsx"),
        ]
        for path in candidates:
            if path and os.path.exists(path):
                return os.path.abspath(path)

        # Glob search for any matching file in current and parent directories
        patterns = [
            "*Regulatory_Reference_Database*.xlsx",
            "../*Regulatory_Reference_Database*.xlsx",
            "step_esg_pipeline/*Regulatory_Reference_Database*.xlsx",
        ]
        for pat in patterns:
            matches = glob.glob(pat)
            if matches:
                return os.path.abspath(matches[0])

        return None

    def load(self, file_path: Optional[str] = None) -> None:
        """
        Loads and validates the regulatory reference workbook.
        Raises RegulatoryDatabaseError with clear diagnostics if missing or invalid.
        """
        if file_path:
            target_path = file_path
        elif self.file_path:
            target_path = self.file_path
        else:
            target_path = self.find_default_file()

        if not target_path or not os.path.exists(target_path):
            raise RegulatoryDatabaseError(
                f"Regulatory Reference Database Excel file not found! Attempted path: '{target_path}'. "
                f"Please ensure 'India_Regulatory_Reference_Database_Refined (1).xlsx' is present in the workspace."
            )

        self.file_path = os.path.abspath(target_path)
        logger.info(f"Loading Regulatory Reference Database from: {self.file_path}")

        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
        except Exception as e:
            raise RegulatoryDatabaseError(f"Failed to read Excel file at '{self.file_path}': {e}") from e

        # Validate required sheets
        sheet_names = wb.sheetnames
        for req_sheet in self.REQUIRED_SHEETS:
            if req_sheet not in sheet_names:
                raise RegulatoryDatabaseError(
                    f"Required sheet '{req_sheet}' is missing from Excel file '{os.path.basename(self.file_path)}'. "
                    f"Available sheets: {sheet_names}"
                )

        # 1. Parse Authorities
        self._parse_authorities(wb["Authorities"])

        # 2. Parse Resources
        self._parse_resources(wb["Resources"])

        # 3. Parse Resource Versions
        self._parse_resource_versions(wb["Resource_Versions"])

        wb.close()
        self.loaded = True
        logger.info(
            f"Successfully loaded Regulatory Reference Database: {len(self.resources)} resources, "
            f"{len(self.authorities)} authorities."
        )

    def _validate_headers(self, sheet, sheet_name: str) -> Dict[str, int]:
        first_row = [cell.value for cell in sheet[1]]
        header_map = {}
        for idx, val in enumerate(first_row):
            if val is not None:
                header_map[str(val).strip()] = idx

        for req_col in self.REQUIRED_COLUMNS.get(sheet_name, []):
            if req_col not in header_map:
                raise RegulatoryDatabaseError(
                    f"Required column '{req_col}' is missing in sheet '{sheet_name}'. "
                    f"Found columns: {list(header_map.keys())}"
                )
        return header_map

    def _parse_authorities(self, sheet) -> None:
        header_map = self._validate_headers(sheet, "Authorities")
        self.authorities.clear()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            auth_id = str(row[header_map["Authority ID"]] or "").strip()
            if not auth_id:
                continue

            country = str(row[header_map["Country"]] or "").strip()
            auth_name = str(row[header_map["Authority"]] or "").strip()
            auth_type = str(row[header_map.get("Authority Type", 0)] or "").strip() if "Authority Type" in header_map else ""
            official_domain = str(row[header_map["Official Domain"]] or "").strip()
            official_web = str(row[header_map.get("Official Website / Portal", 0)] or "").strip() if "Official Website / Portal" in header_map else ""
            trust_level = str(row[header_map.get("Trust Level", 0)] or "Official").strip() if "Trust Level" in header_map else "Official"
            notes = str(row[header_map.get("Notes", 0)] or "").strip() if "Notes" in header_map else ""

            self.authorities[auth_id] = RegulatoryAuthority(
                authority_id=auth_id,
                country=country,
                authority_name=auth_name,
                authority_type=auth_type,
                official_domain=official_domain,
                official_website=official_web,
                trust_level=trust_level,
                notes=notes,
            )

    def _parse_resources(self, sheet) -> None:
        header_map = self._validate_headers(sheet, "Resources")
        self.resources.clear()
        self.resources_by_country.clear()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            res_id = str(row[header_map["Resource ID"]] or "").strip()
            if not res_id:
                continue

            country = str(row[header_map["Country"]] or "").strip()
            name = str(row[header_map["Resource Name / Title"]] or "").strip()
            issuing_auth = str(row[header_map["Issuing Authority"]] or "").strip()
            auth_type = str(row[header_map.get("Authority Type", 0)] or "").strip() if "Authority Type" in header_map else ""
            topic = str(row[header_map.get("Regulatory Topic / Area", 0)] or "").strip() if "Regulatory Topic / Area" in header_map else ""
            doc_type = str(row[header_map.get("Document Type", 0)] or "").strip() if "Document Type" in header_map else ""
            official_domain = str(row[header_map["Official Domain"]] or "").strip()
            official_url = str(row[header_map["Official URL"]] or "").strip()
            pub_date = str(row[header_map.get("Publication / Enactment Date", 0)] or "").strip() if "Publication / Enactment Date" in header_map else ""
            status_ver = str(row[header_map.get("Status / Verification State", 0)] or "").strip() if "Status / Verification State" in header_map else ""
            lifecycle = str(row[header_map.get("Lifecycle Notes / Currentness", 0)] or "").strip() if "Lifecycle Notes / Currentness" in header_map else ""
            related_url = str(row[header_map.get("Related / Updated URL", 0)] or "").strip() if "Related / Updated URL" in header_map else ""
            val_rule = str(row[header_map.get("Validator Rule / Guidance", 0)] or "").strip() if "Validator Rule / Guidance" in header_map else ""
            impl_notes = str(row[header_map.get("Implementation Notes", 0)] or "").strip() if "Implementation Notes" in header_map else ""

            # Match authority details from authorities sheet
            matched_auth = None
            for auth in self.authorities.values():
                if (auth.country.lower() == country.lower() and
                    (auth.authority_name.lower() in issuing_auth.lower() or
                     issuing_auth.lower() in auth.authority_name.lower() or
                     auth.official_domain.lower() in official_domain.lower())):
                    matched_auth = auth
                    break

            res = RegulatoryResource(
                resource_id=res_id,
                country=country,
                name=name,
                issuing_authority=issuing_auth,
                authority_type=auth_type,
                regulatory_topic=topic,
                document_type=doc_type,
                official_domain=official_domain,
                official_url=official_url,
                publication_date=pub_date,
                status_verification_state=status_ver,
                lifecycle_notes=lifecycle,
                related_updated_url=related_url,
                validator_rule=val_rule,
                implementation_notes=impl_notes,
                authority_details=matched_auth,
            )
            self.resources[res_id] = res

            c_key = country.lower()
            if c_key not in self.resources_by_country:
                self.resources_by_country[c_key] = []
            self.resources_by_country[c_key].append(res)

    def _parse_resource_versions(self, sheet) -> None:
        header_map = self._validate_headers(sheet, "Resource_Versions")

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            res_id = str(row[header_map["Resource ID"]] or "").strip()
            if not res_id or res_id not in self.resources:
                continue

            update_name = str(row[header_map["Resource / Update"]] or "").strip()
            relationship = str(row[header_map["Relationship"]] or "").strip()
            issuer = str(row[header_map.get("Issuer", 0)] or "").strip() if "Issuer" in header_map else ""
            date_str = str(row[header_map.get("Date", 0)] or "").strip() if "Date" in header_map else ""
            official_url = str(row[header_map.get("Official URL", 0)] or "").strip() if "Official URL" in header_map else ""
            val_treat = str(row[header_map.get("Validator Treatment", 0)] or "").strip() if "Validator Treatment" in header_map else ""
            notes = str(row[header_map.get("Notes", 0)] or "").strip() if "Notes" in header_map else ""

            ver = ResourceVersion(
                resource_id=res_id,
                resource_update=update_name,
                issuer=issuer,
                date=date_str,
                relationship=relationship,
                official_url=official_url,
                validator_treatment=val_treat,
                notes=notes,
            )
            self.resources[res_id].versions.append(ver)

    def get_resource(self, resource_id: str) -> Optional[RegulatoryResource]:
        """Returns the resource by ID (e.g. 'IND-001')."""
        return self.resources.get(resource_id)

    def match_resource(self, jurisdiction: str, url: str = "", text: str = "") -> Optional[RegulatoryResource]:
        """
        Matches a STEP link to its canonical RegulatoryResource in the Excel database.
        Checks URL matching, official domains, and keyword heuristics.
        """
        if not self.loaded:
            try:
                self.load()
            except Exception:
                return None

        # Filter by country (case-insensitive)
        candidates = self.resources_by_country.get(jurisdiction.lower(), [])
        if not candidates and (jurisdiction.lower() in ("india", "in") or "india" in text.lower()):
            candidates = self.resources_by_country.get("india", [])

        if not candidates:
            return None

        clean_url = (url or "").lower().strip()
        parsed_url = urlparse(clean_url)
        netloc = parsed_url.netloc.lower()
        path = parsed_url.path.lower()
        text_low = (text or "").lower()

        # 1. Exact or Substring URL Match
        for res in candidates:
            all_urls = [u.lower().strip() for u in res.get_related_urls()]
            for u in all_urls:
                if u == clean_url:
                    return res
                # Match core circular/act path
                parsed_u = urlparse(u)
                if parsed_u.path and len(parsed_u.path) > 10 and parsed_u.path in path:
                    return res

        # 2. Resource-Specific Heuristics for India Pilot
        for res in candidates:
            rid = res.resource_id
            if rid == "IND-001":
                # SEBI BRSR Circular
                if "sebi.gov.in" in netloc and ("50096" in path or "business-responsibility" in path or "brsr" in path):
                    return res
                if "brsr" in text_low or "business responsibility and sustainability reporting" in text_low:
                    return res
            elif rid == "IND-002":
                # Companies Act 2013 Section 135 & CSR Rules
                if ("mca.gov.in" in netloc or "indiacode.nic.in" in netloc) and ("companiesact2013" in path or "csr" in path or "123456789/2114" in path):
                    return res
                if ("companies act" in text_low and "135" in text_low) or ("csr rules" in text_low):
                    return res
            elif rid == "IND-003":
                # RBI Environmental and Social Risk Management Guidelines
                if "rbi.org.in" in netloc and ("12213" in path or "4393" in path or "notificationuser" in path):
                    return res
                if ("rbi" in text_low or "reserve bank" in text_low) and ("environmental and social risk" in text_low or "sustainable finance" in text_low):
                    return res

        # 3. Fallback Title/Text Substring Match
        for res in candidates:
            r_name_low = res.name.lower()
            if r_name_low in text_low:
                return res

        return None
