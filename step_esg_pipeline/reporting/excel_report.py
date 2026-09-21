"""
Excel report generator for the STEP ESG Link Monitoring Pipeline.
Produces a formatted 5-sheet .xlsx file for human review:
1. 'STEP ESG Link Audit'
2. 'Evidence'
3. 'Candidates'
4. 'Country Summary'
5. 'Manual Review'
Requires: openpyxl >= 3.1.0
"""
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Dict, Any
from urllib.parse import urlparse

from database.models import LinkRecord

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# Color palette (Modern ESG executive reporting theme)
COLOR_HEADER_BG   = "1E3A5F"   # Deep Navy
COLOR_HEADER_FG   = "FFFFFF"
COLOR_KEEP        = "D6F5E3"   # Soft Green
COLOR_UPDATE      = "E3F2FD"   # Soft Blue
COLOR_REPLACE     = "FFF3CD"   # Soft Amber
COLOR_FIX_BROKEN  = "FFE0B2"   # Soft Orange
COLOR_REVIEW      = "FDECEA"   # Soft Pink/Red
COLOR_ALT_ROW     = "F8FAFC"   # Light Blue-Grey
COLOR_SUBHEADER   = "2C3E50"

DECISION_COLORS = {
    "KEEP": COLOR_KEEP,
    "UPDATE": COLOR_UPDATE,
    "REPLACE": COLOR_REPLACE,
    "FIX_BROKEN_LINK": COLOR_FIX_BROKEN,
    "MANUAL_REVIEW": COLOR_REVIEW,
    "HUMAN_REVIEW_REQUIRED": COLOR_REVIEW,
}


def _decision_fill(decision: str) -> PatternFill:
    color = DECISION_COLORS.get((decision or "").upper(), "FFFFFF")
    return PatternFill("solid", fgColor=color)


AUDIT_COLUMNS = [
    ("No.", lambda r, i: i),
    ("Section", lambda r, i: r.section or "ESG Legislative Landscape"),
    ("Country", lambda r, i: r.country or r.jurisdiction or ""),
    ("Country Conf", lambda r, i: f"{int((r.country_confidence or 1.0) * 100)}%"),
    ("Instrument Name", lambda r, i: r.instrument_name or r.step_description or ""),
    ("Instrument Type", lambda r, i: r.instrument_type or ""),
    ("Regulatory Topic", lambda r, i: r.regulatory_topic or r.topic or ""),
    ("Regulated Population", lambda r, i: r.regulated_population or ""),
    ("Original URL", lambda r, i: r.original_url or ""),
    ("Final URL", lambda r, i: r.final_url or ""),
    ("Technical Status", lambda r, i: r.technical_status or ("LIVE" if r.http_status == 200 else "BROKEN")),
    ("Technical Page Title", lambda r, i: r.technical_page_title or ""),
    ("Authority", lambda r, i: r.authority or r.source_organisation or ""),
    ("Authority Status", lambda r, i: r.authority_status or ""),
    ("Regulatory Lifecycle", lambda r, i: r.regulatory_status or ""),
    ("Candidate URL", lambda r, i: r.replacement_url or ""),
    ("Candidate Title", lambda r, i: r.replacement_title or ""),
    ("Candidate Authority", lambda r, i: r.candidate_authority or ""),
    ("Candidate Status", lambda r, i: r.candidate_status or ("REPLACEMENT_VERIFIED" if r.replacement_verified else "NONE_FOUND")),
    ("Replacement Relationship", lambda r, i: r.replacement_relationship or ""),
    ("Tech Conf", lambda r, i: f"{int((r.technical_confidence or 0.0) * 100)}%"),
    ("Auth Conf", lambda r, i: f"{int((r.authority_confidence or 0.0) * 100)}%"),
    ("ID Conf", lambda r, i: f"{int((r.identity_confidence or 0.0) * 100)}%"),
    ("Curr Conf", lambda r, i: f"{int((r.currentness_confidence or 0.0) * 100)}%"),
    ("Rep Conf", lambda r, i: f"{int((r.replacement_confidence or 0.0) * 100)}%"),
    ("Overall Conf", lambda r, i: f"{int(((r.overall_confidence or r.confidence_score) or 0.0) * 100)}%"),
    ("Final Decision", lambda r, i: r.final_decision or r.recommended_action or "MANUAL_REVIEW"),
    ("Why Summary", lambda r, i: r.why_summary or ""),
]


class ExcelReport:
    @staticmethod
    def generate(records: List[LinkRecord], path: str) -> None:
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl>=3.1.0")

        wb = openpyxl.Workbook()

        # ----------------------------------------------------------------
        # Sheet 1: STEP ESG Link Audit (Main Audit Sheet)
        # ----------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "STEP ESG Link Audit"
        ExcelReport._write_audit_sheet(ws1, records)

        # ----------------------------------------------------------------
        # Sheet 2: Evidence (Category-tagged evidence records)
        # ----------------------------------------------------------------
        ws2 = wb.create_sheet("Evidence")
        ExcelReport._write_evidence_sheet(ws2, records)

        # ----------------------------------------------------------------
        # Sheet 3: Candidates (Replacements and official discoveries)
        # ----------------------------------------------------------------
        ws3 = wb.create_sheet("Candidates")
        ExcelReport._write_candidates_sheet(ws3, records)

        # ----------------------------------------------------------------
        # Sheet 4: Country Summary (Breakdown by 13 jurisdictions)
        # ----------------------------------------------------------------
        ws4 = wb.create_sheet("Country Summary")
        ExcelReport._write_country_summary(ws4, records)

        # ----------------------------------------------------------------
        # Sheet 5: Manual Review (Items requiring human verification)
        # ----------------------------------------------------------------
        ws5 = wb.create_sheet("Manual Review")
        ExcelReport._write_manual_review_sheet(ws5, records)

        wb.save(path)

    @staticmethod
    def _write_audit_sheet(ws, records: List[LinkRecord]) -> None:
        num_cols = len(AUDIT_COLUMNS)
        # Title row
        ws.merge_cells(f"A1:{get_column_letter(num_cols)}1")
        title_cell = ws["A1"]
        title_cell.value = "STEP ESG Legislative Landscape — Comprehensive Regulatory Link Audit"
        title_cell.font = Font(name="Calibri", bold=True, size=14, color=COLOR_HEADER_FG)
        title_cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Subtitle row
        ws.merge_cells(f"A2:{get_column_letter(num_cols)}2")
        ts_cell = ws["A2"]
        ts_cell.value = f"Audit Timestamp: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | Scoped Links: {len(records)} | Section: ESG Legislative Landscape"
        ts_cell.font = Font(name="Calibri", italic=True, size=10, color="666666")
        ts_cell.alignment = Alignment(horizontal="center")

        # Header row
        header_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        header_font = Font(name="Calibri", bold=True, size=10, color=COLOR_HEADER_FG)
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, bottom=thin)

        for col_idx, (col_label, _) in enumerate(AUDIT_COLUMNS, start=1):
            cell = ws.cell(row=3, column=col_idx, value=col_label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[3].height = 36

        # Data rows
        normal_font = Font(name="Calibri", size=9)
        url_font = Font(name="Calibri", size=9, color="1565C0", underline="single")
        center = Alignment(horizontal="center", vertical="top", wrap_text=True)
        left = Alignment(horizontal="left", vertical="top", wrap_text=True)

        for row_idx, record in enumerate(records, start=1):
            excel_row = row_idx + 3
            decision = record.final_decision or record.recommended_action or "MANUAL_REVIEW"
            row_fill = _decision_fill(decision) if decision != "KEEP" else (
                PatternFill("solid", fgColor=COLOR_ALT_ROW if row_idx % 2 == 0 else "FFFFFF")
            )

            for col_idx, (col_label, extractor) in enumerate(AUDIT_COLUMNS, start=1):
                value = extractor(record, row_idx)
                cell = ws.cell(row=excel_row, column=col_idx, value=value)
                cell.fill = row_fill
                cell.border = border

                if col_label in ("Original URL", "Final URL", "Candidate URL") and value and str(value).startswith("http"):
                    cell.font = url_font
                    try:
                        cell.hyperlink = str(value)
                    except Exception:
                        pass
                    cell.alignment = left
                elif col_label in ("No.", "Country Conf", "Tech Conf", "Auth Conf", "ID Conf", "Curr Conf", "Rep Conf", "Overall Conf", "Final Decision"):
                    cell.font = normal_font
                    cell.alignment = center
                else:
                    cell.font = normal_font
                    cell.alignment = left

        # Column widths
        for col_idx in range(1, num_cols + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 24
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["E"].width = 38
        ws.column_dimensions["I"].width = 45
        ws.column_dimensions["J"].width = 45
        ws.column_dimensions["P"].width = 45
        ws.column_dimensions["AB"].width = 55

        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(num_cols)}3"

    @staticmethod
    def _write_evidence_sheet(ws, records: List[LinkRecord]) -> None:
        headers = ["No.", "Country", "Instrument Name", "Category", "Finding / Evidence Detail"]
        ws.merge_cells("A1:E1")
        ws["A1"].value = "STEP ESG Link Audit — Evidence Trail"
        ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=COLOR_HEADER_FG)
        ws["A1"].fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 26

        header_fill = PatternFill("solid", fgColor=COLOR_SUBHEADER)
        header_font = Font(name="Calibri", bold=True, size=10, color=COLOR_HEADER_FG)
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 24

        curr_row = 3
        thin = Side(style="thin", color="E0E0E0")
        border = Border(left=thin, right=thin, bottom=thin)

        for rec_idx, r in enumerate(records, start=1):
            evidence_items = []
            ev_json = getattr(r, "evidence_json", "")
            if ev_json and ev_json != "[]":
                try:
                    evidence_items = json.loads(ev_json)
                except Exception:
                    pass
            if not evidence_items and r.evidence and r.evidence != "[]":
                try:
                    raw_list = json.loads(r.evidence)
                    evidence_items = [{"type": "general", "finding": str(x)} for x in raw_list]
                except Exception:
                    pass

            country = r.country or r.jurisdiction or ""
            inst = r.instrument_name or r.step_description or ""

            for it in evidence_items:
                ev_type = (it.get("type") or "general").upper()
                finding = it.get("finding") or it.get("detail") or ""
                ws.cell(row=curr_row, column=1, value=rec_idx).alignment = Alignment(horizontal="center")
                ws.cell(row=curr_row, column=2, value=country)
                ws.cell(row=curr_row, column=3, value=inst)
                ws.cell(row=curr_row, column=4, value=ev_type).alignment = Alignment(horizontal="center")
                ws.cell(row=curr_row, column=5, value=finding)

                for c in range(1, 6):
                    ws.cell(row=curr_row, column=c).border = border
                curr_row += 1

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 75
        ws.freeze_panes = "A3"

    @staticmethod
    def _write_candidates_sheet(ws, records: List[LinkRecord]) -> None:
        headers = [
            "No.", "Country", "Original Instrument", "Candidate URL", "Candidate Title",
            "Candidate Authority", "Candidate Status", "Replacement Relationship"
        ]
        ws.merge_cells("A1:H1")
        ws["A1"].value = "STEP ESG Link Audit — Replacement & Official Source Candidates"
        ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=COLOR_HEADER_FG)
        ws["A1"].fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 26

        header_fill = PatternFill("solid", fgColor=COLOR_SUBHEADER)
        header_font = Font(name="Calibri", bold=True, size=10, color=COLOR_HEADER_FG)
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 24

        curr_row = 3
        thin = Side(style="thin", color="E0E0E0")
        border = Border(left=thin, right=thin, bottom=thin)

        for rec_idx, r in enumerate(records, start=1):
            if not r.replacement_url and r.candidate_status in ("", "NONE_FOUND", "NO_REPLACEMENT_REQUIRED"):
                continue

            ws.cell(row=curr_row, column=1, value=rec_idx).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=2, value=r.country or r.jurisdiction or "")
            ws.cell(row=curr_row, column=3, value=r.instrument_name or r.step_description or "")
            ws.cell(row=curr_row, column=4, value=r.replacement_url or "")
            ws.cell(row=curr_row, column=5, value=r.replacement_title or "")
            ws.cell(row=curr_row, column=6, value=r.candidate_authority or "")
            ws.cell(row=curr_row, column=7, value=r.candidate_status or "").alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=8, value=r.replacement_relationship or "").alignment = Alignment(horizontal="center")

            for c in range(1, 9):
                ws.cell(row=curr_row, column=c).border = border
            curr_row += 1

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 50
        ws.column_dimensions["E"].width = 40
        ws.column_dimensions["F"].width = 25
        ws.column_dimensions["G"].width = 24
        ws.column_dimensions["H"].width = 28
        ws.freeze_panes = "A3"

    @staticmethod
    def _write_country_summary(ws, records: List[LinkRecord]) -> None:
        headers = [
            "Country", "Total Links", "Live / Valid", "Blocked / Challenge", "Broken / 404",
            "Tier 1 Official", "Tier 2 Exchange", "Tier 3 Aggregator",
            "Current in Force", "Superseded/Amended", "Consultation/Draft", "Unknown Lifecycle",
            "KEEP", "UPDATE", "REPLACE", "FIX_BROKEN_LINK", "MANUAL_REVIEW"
        ]
        ws.merge_cells("A1:Q1")
        ws["A1"].value = "STEP ESG Link Audit — Country Level Executive Summary"
        ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=COLOR_HEADER_FG)
        ws["A1"].fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 26

        header_fill = PatternFill("solid", fgColor=COLOR_SUBHEADER)
        header_font = Font(name="Calibri", bold=True, size=9, color=COLOR_HEADER_FG)
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[2].height = 28

        # Group by country
        by_country: Dict[str, List[LinkRecord]] = {}
        for r in records:
            c = r.country or r.jurisdiction or "UNKNOWN"
            by_country.setdefault(c, []).append(r)

        curr_row = 3
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, bottom=thin)

        for country, items in sorted(by_country.items()):
            tot = len(items)
            live = sum(1 for r in items if r.technical_status in ("LIVE", "PDF_VALID", "ACCESSIBLE_VIA_BROWSER") or r.http_status == 200)
            blocked = sum(1 for r in items if r.technical_status == "ACCESS_BLOCKED" or r.http_status in (403, 202))
            broken = sum(1 for r in items if r.technical_status in ("HTTP_404", "BROKEN", "SOFT_404") or r.http_status in (404, 410))

            t1 = sum(1 for r in items if "TIER_1" in (r.authority_status or "").upper())
            t2 = sum(1 for r in items if "TIER_2" in (r.authority_status or "").upper())
            t3 = sum(1 for r in items if "TIER_3" in (r.authority_status or "").upper())

            cur = sum(1 for r in items if (r.regulatory_status or "").upper() in ("CURRENT_IN_FORCE", "LEGALLY_BINDING_IN_FORCE"))
            sup = sum(1 for r in items if (r.regulatory_status or "").upper() in ("SUPERSEDED", "REPEALED", "AMENDED"))
            con = sum(1 for r in items if (r.regulatory_status or "").upper() in ("CONSULTATION", "CONSULTATION_DRAFT"))
            unk = sum(1 for r in items if (r.regulatory_status or "").upper() in ("UNKNOWN", ""))

            keep = sum(1 for r in items if (r.final_decision or r.recommended_action or "").upper() == "KEEP")
            upd = sum(1 for r in items if (r.final_decision or r.recommended_action or "").upper() == "UPDATE")
            rep = sum(1 for r in items if (r.final_decision or r.recommended_action or "").upper() == "REPLACE")
            fix = sum(1 for r in items if (r.final_decision or r.recommended_action or "").upper() == "FIX_BROKEN_LINK")
            rev = sum(1 for r in items if (r.final_decision or r.recommended_action or "").upper() in ("MANUAL_REVIEW", "HUMAN_REVIEW_REQUIRED"))

            vals = [country, tot, live, blocked, broken, t1, t2, t3, cur, sup, con, unk, keep, upd, rep, fix, rev]
            for col_idx, v in enumerate(vals, start=1):
                cell = ws.cell(row=curr_row, column=col_idx, value=v)
                cell.border = border
                cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "center")
            curr_row += 1

        for c in range(1, 18):
            ws.column_dimensions[get_column_letter(c)].width = 16
        ws.column_dimensions["A"].width = 24
        ws.freeze_panes = "B3"

    @staticmethod
    def _write_manual_review_sheet(ws, records: List[LinkRecord]) -> None:
        headers = [
            "No.", "Country", "Instrument Name", "Original URL", "Technical Status",
            "Regulatory Status", "Overall Confidence", "Primary Review Reason", "Why Summary"
        ]
        ws.merge_cells("A1:I1")
        ws["A1"].value = "STEP ESG Link Audit — Actionable Manual Review Queue"
        ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=COLOR_HEADER_FG)
        ws["A1"].fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 26

        header_fill = PatternFill("solid", fgColor=COLOR_SUBHEADER)
        header_font = Font(name="Calibri", bold=True, size=10, color=COLOR_HEADER_FG)
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 24

        curr_row = 3
        thin = Side(style="thin", color="E0E0E0")
        border = Border(left=thin, right=thin, bottom=thin)

        review_records = [
            r for r in records
            if (r.final_decision or r.recommended_action or "").upper() in ("MANUAL_REVIEW", "HUMAN_REVIEW_REQUIRED")
            or r.human_review_required
        ]

        for idx, r in enumerate(review_records, start=1):
            # Determine primary review reason
            if r.http_status in (403, 202) or r.technical_status in ("ACCESS_BLOCKED", "WAF_CHALLENGE"):
                reason = "Access Blocked / Bot Challenge Unresolved"
            elif (r.regulatory_status or "").upper() == "UNKNOWN":
                reason = "Regulatory Lifecycle Status Unknown (Affirmative proof required)"
            elif (r.regulatory_status or "").upper() in ("CONSULTATION", "CONSULTATION_DRAFT"):
                reason = "Consultation / Draft paper without confirmed final regulation"
            elif r.replacement_required and not r.replacement_verified:
                reason = "Replacement Required but Candidate Unverified"
            elif r.confidence_score < 0.65 or r.overall_confidence < 0.65:
                reason = "Confidence Score Below Verification Threshold"
            else:
                reason = r.issue_description or r.regulatory_status_reason or "Review required"

            ws.cell(row=curr_row, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=2, value=r.country or r.jurisdiction or "")
            ws.cell(row=curr_row, column=3, value=r.instrument_name or r.step_description or "")
            ws.cell(row=curr_row, column=4, value=r.original_url or "")
            ws.cell(row=curr_row, column=5, value=r.technical_status or "").alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=6, value=r.regulatory_status or "").alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=7, value=f"{int(((r.overall_confidence or r.confidence_score) or 0.0) * 100)}%").alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=8, value=reason)
            ws.cell(row=curr_row, column=9, value=r.why_summary or "")

            for c in range(1, 10):
                ws.cell(row=curr_row, column=c).border = border
            curr_row += 1

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 45
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 22
        ws.column_dimensions["G"].width = 18
        ws.column_dimensions["H"].width = 38
        ws.column_dimensions["I"].width = 50
        ws.freeze_panes = "A3"
